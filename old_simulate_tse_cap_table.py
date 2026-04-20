#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


@dataclass(frozen=True)
class Instrument:
    layer: str
    rank: int
    priority_label: str
    name: str
    principal: float
    maturity: str
    interest_type: str


@dataclass(frozen=True)
class CapTableModel:
    workbook_name: str
    instruments: list[Instrument]
    base_ebitda: float
    multiple_low: float
    multiple_mid: float
    multiple_high: float
    listed_debt: float
    other_claims: float
    total_debt_to_clear: float


@dataclass(frozen=True)
class SimulationAssumptions:
    runs: int
    seed: int
    ebitda_cv: float
    debt_volatility: float


LAYER_ORDER = {
    "Super": 1,
    "1L": 2,
    "2L": 3,
    "Junior": 4,
    "Equity": 5,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run a Monte Carlo recovery analysis on the TSE cap table workbook."
    )
    parser.add_argument(
        "--input",
        default="TSE Cap Table .xlsx",
        help="Path to the cap table workbook.",
    )
    parser.add_argument(
        "--output-dir",
        default="output",
        help="Directory for generated CSV, JSON, and HTML artifacts.",
    )
    parser.add_argument(
        "--runs",
        type=int,
        default=50000,
        help="Monte Carlo iteration count.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=7,
        help="Random seed for reproducibility.",
    )
    parser.add_argument(
        "--ebitda-cv",
        type=float,
        default=0.22,
        help="Coefficient of variation for the EBITDA lognormal distribution.",
    )
    parser.add_argument(
        "--debt-volatility",
        type=float,
        default=0.03,
        help="Standard deviation for the debt-balance shock applied pro rata to all claims.",
    )
    return parser.parse_args()


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.split()).strip()
    return str(value).strip()


def to_float(value: Any) -> float:
    if value in (None, "", "#DIV/0!"):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def canonicalize_layer(priority_label: str) -> tuple[str, int]:
    normalized = priority_label.lower()
    if "super" in normalized:
        return "Super", LAYER_ORDER["Super"]
    if "1st" in normalized or "senior secured" in normalized:
        return "1L", LAYER_ORDER["1L"]
    if "2nd" in normalized or "2nd lien" in normalized:
        return "2L", LAYER_ORDER["2L"]
    if "junior" in normalized:
        return "Junior", LAYER_ORDER["Junior"]
    if "equity" in normalized:
        return "Equity", LAYER_ORDER["Equity"]
    raise ValueError(f"Unsupported priority label: {priority_label!r}")


def maybe_layer(priority_label: str) -> tuple[str, int] | None:
    try:
        return canonicalize_layer(priority_label)
    except ValueError:
        return None


def load_cap_table(workbook_path: Path) -> CapTableModel:
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    instruments: list[Instrument] = []
    base_ebitda = 0.0
    multiple_low = 0.0
    multiple_mid = 0.0
    multiple_high = 0.0
    total_debt_to_clear = 0.0
    other_claims = 0.0
    in_stack_section = True

    for row in ws.iter_rows(values_only=True):
        priority = normalize_text(row[0])
        raw_instrument_name = row[1]
        instrument_name = normalize_text(raw_instrument_name)
        principal = to_float(row[2])
        maturity = normalize_text(row[3])
        interest_type = normalize_text(row[4])

        parsed_layer = maybe_layer(priority)
        if instrument_name == "Total Debt to Clear" or priority == "Adjusted EBITDA":
            in_stack_section = False

        if (
            in_stack_section
            and instrument_name
            and isinstance(raw_instrument_name, str)
            and "total" not in instrument_name.lower()
            and parsed_layer
        ):
            layer, rank = parsed_layer
            instruments.append(
                Instrument(
                    layer=layer,
                    rank=rank,
                    priority_label=priority,
                    name=instrument_name,
                    principal=principal,
                    maturity=maturity or "n/a",
                    interest_type=interest_type or "n/a",
                )
            )

        if priority == "Adjusted EBITDA":
            base_ebitda = to_float(row[1])
        elif priority == "Low (distressed case)":
            multiple_low = to_float(row[1])
        elif priority == "Mid":
            multiple_mid = to_float(row[1])
        elif priority == "High (healthy case)":
            multiple_high = to_float(row[1])
        elif instrument_name == "Total Debt to Clear":
            total_debt_to_clear = to_float(row[2])
        elif priority == "Other":
            other_claims = to_float(row[2])

    if not instruments:
        raise ValueError("No instruments were found in the workbook.")

    listed_debt = sum(item.principal for item in instruments)

    return CapTableModel(
        workbook_name=workbook_path.name,
        instruments=sorted(instruments, key=lambda item: (item.rank, item.name)),
        base_ebitda=base_ebitda,
        multiple_low=multiple_low,
        multiple_mid=multiple_mid,
        multiple_high=multiple_high,
        listed_debt=listed_debt,
        other_claims=other_claims,
        total_debt_to_clear=total_debt_to_clear or (listed_debt + other_claims),
    )


def lognormal_samples(
    mean: float,
    coefficient_of_variation: float,
    size: int,
    rng: np.random.Generator,
) -> np.ndarray:
    sigma2 = math.log(1 + coefficient_of_variation**2)
    sigma = math.sqrt(sigma2)
    mu = math.log(mean) - sigma2 / 2
    return rng.lognormal(mean=mu, sigma=sigma, size=size)


def summarize_distribution(values: np.ndarray) -> dict[str, float]:
    return {
        "mean": float(np.mean(values)),
        "p10": float(np.percentile(values, 10)),
        "p25": float(np.percentile(values, 25)),
        "p50": float(np.percentile(values, 50)),
        "p75": float(np.percentile(values, 75)),
        "p90": float(np.percentile(values, 90)),
    }


def build_histogram(values: np.ndarray, bins: int) -> list[dict[str, float]]:
    counts, edges = np.histogram(values, bins=bins)
    output: list[dict[str, float]] = []
    for index, count in enumerate(counts):
        output.append(
            {
                "start": float(edges[index]),
                "end": float(edges[index + 1]),
                "count": int(count),
            }
        )
    return output


def run_waterfall(
    layer_names: list[str],
    layer_balances: dict[str, np.ndarray],
    enterprise_value: np.ndarray,
    other_claims: np.ndarray,
) -> tuple[dict[str, np.ndarray], dict[str, np.ndarray], np.ndarray]:
    remaining_value = enterprise_value.copy()
    recovered_amounts: dict[str, np.ndarray] = {}
    recovery_rates: dict[str, np.ndarray] = {}

    for layer in layer_names:
        balance = layer_balances[layer]
        recovered = np.minimum(balance, remaining_value)
        recovered_amounts[layer] = recovered
        recovery_rates[layer] = np.divide(
            recovered,
            balance,
            out=np.zeros_like(recovered),
            where=balance > 0,
        )
        remaining_value = np.maximum(remaining_value - recovered, 0)

    equity_value = np.maximum(remaining_value - other_claims, 0)
    return recovered_amounts, recovery_rates, equity_value


def scenario_waterfall(ev: float, layer_totals: dict[str, float], other_claims: float) -> list[dict[str, float | str]]:
    remaining = ev
    rows: list[dict[str, float | str]] = []
    for layer in ("Super", "1L", "2L"):
        debt = layer_totals.get(layer, 0.0)
        recovered = min(remaining, debt)
        recovery_pct = recovered / debt if debt else 0.0
        rows.append(
            {
                "layer": layer,
                "claim": debt,
                "recovered": recovered,
                "recovery_pct": recovery_pct,
            }
        )
        remaining = max(remaining - recovered, 0.0)

    reserve_recovered = min(remaining, other_claims)
    rows.append(
        {
            "layer": "Other claims reserve",
            "claim": other_claims,
            "recovered": reserve_recovered,
            "recovery_pct": reserve_recovered / other_claims if other_claims else 0.0,
        }
    )
    remaining = max(remaining - other_claims, 0.0)
    rows.append(
        {
            "layer": "Equity residual",
            "claim": 0.0,
            "recovered": remaining,
            "recovery_pct": 0.0,
        }
    )
    return rows


def build_sensitivity_matrix(model: CapTableModel) -> dict[str, Any]:
    ebitda_values = [model.base_ebitda * factor for factor in (0.75, 0.9, 1.0, 1.1, 1.25)]
    multiple_values = [
        model.multiple_low,
        statistics.mean([model.multiple_low, model.multiple_mid]),
        model.multiple_mid,
        statistics.mean([model.multiple_mid, model.multiple_high]),
        model.multiple_high,
    ]
    layer_totals: dict[str, float] = {}
    for instrument in model.instruments:
        layer_totals[instrument.layer] = layer_totals.get(instrument.layer, 0.0) + instrument.principal

    cells = []
    max_equity = 0.0
    for ebitda in ebitda_values:
        row = []
        for multiple in multiple_values:
            ev = ebitda * multiple
            remaining = ev
            for layer in ("Super", "1L", "2L"):
                remaining = max(remaining - layer_totals.get(layer, 0.0), 0.0)
            equity = max(remaining - model.other_claims, 0.0)
            second_lien_recovered = min(
                max(ev - layer_totals.get("Super", 0.0) - layer_totals.get("1L", 0.0), 0.0),
                layer_totals.get("2L", 0.0),
            )
            second_lien_recovery = (
                second_lien_recovered / layer_totals.get("2L", 1.0)
                if layer_totals.get("2L", 0.0)
                else 0.0
            )
            row.append(
                {
                    "enterprise_value": ev,
                    "equity_value": equity,
                    "second_lien_recovery": second_lien_recovery,
                }
            )
            max_equity = max(max_equity, equity)
        cells.append(row)

    return {
        "ebitda_values": ebitda_values,
        "multiple_values": multiple_values,
        "cells": cells,
        "max_equity": max_equity,
    }


def simulate(model: CapTableModel, assumptions: SimulationAssumptions) -> dict[str, Any]:
    rng = np.random.default_rng(assumptions.seed)
    runs = assumptions.runs

    ebitda = lognormal_samples(model.base_ebitda, assumptions.ebitda_cv, runs, rng)
    exit_multiple = rng.triangular(model.multiple_low, model.multiple_mid, model.multiple_high, runs)
    debt_scale = np.clip(
        1 + rng.normal(0, assumptions.debt_volatility, runs),
        0.75,
        1.35,
    )

    enterprise_value = ebitda * exit_multiple
    other_claims = model.other_claims * debt_scale

    layer_totals: dict[str, float] = {}
    for instrument in model.instruments:
        layer_totals[instrument.layer] = layer_totals.get(instrument.layer, 0.0) + instrument.principal

    ordered_layers = [layer for layer in ("Super", "1L", "2L") if layer in layer_totals]
    layer_balances = {layer: layer_totals[layer] * debt_scale for layer in ordered_layers}

    recovered_amounts, recovery_rates, equity_value = run_waterfall(
        ordered_layers,
        layer_balances,
        enterprise_value,
        other_claims,
    )

    total_debt_to_clear = model.total_debt_to_clear * debt_scale
    insolvency_probability = float(np.mean(enterprise_value < total_debt_to_clear))

    layer_stats: list[dict[str, Any]] = []
    for layer in ordered_layers:
        balance = layer_totals[layer]
        rate = recovery_rates[layer]
        recovered = recovered_amounts[layer]
        layer_stats.append(
            {
                "layer": layer,
                "current_claim": balance,
                "expected_claim_at_exit": float(np.mean(layer_balances[layer])),
                "expected_recovered": float(np.mean(recovered)),
                "expected_recovery_pct": float(np.mean(rate)),
                "money_good_probability": float(np.mean(rate >= 0.999)),
                "recovery_distribution": summarize_distribution(rate),
            }
        )

    instrument_stats: list[dict[str, Any]] = []
    layer_rate_lookup = {row["layer"]: recovery_rates[row["layer"]] for row in layer_stats}
    layer_balance_lookup = {layer: layer_balances[layer] for layer in ordered_layers}
    for instrument in model.instruments:
        if instrument.layer not in layer_rate_lookup:
            continue
        claim_share = instrument.principal / layer_totals[instrument.layer]
        recovered = recovered_amounts[instrument.layer] * claim_share
        rate = layer_rate_lookup[instrument.layer]
        instrument_stats.append(
            {
                "layer": instrument.layer,
                "instrument": instrument.name,
                "principal": instrument.principal,
                "maturity": instrument.maturity,
                "interest_type": instrument.interest_type,
                "expected_claim_at_exit": float(np.mean(layer_balance_lookup[instrument.layer] * claim_share)),
                "expected_recovered": float(np.mean(recovered)),
                "expected_recovery_pct": float(np.mean(rate)),
                "money_good_probability": float(np.mean(rate >= 0.999)),
                "recovery_distribution": summarize_distribution(rate),
            }
        )

    low_ev = model.base_ebitda * model.multiple_low
    base_ev = model.base_ebitda * model.multiple_mid
    high_ev = model.base_ebitda * model.multiple_high

    scenarios = [
        {
            "label": "Distressed",
            "enterprise_value": low_ev,
            "waterfall": scenario_waterfall(low_ev, layer_totals, model.other_claims),
        },
        {
            "label": "Base",
            "enterprise_value": base_ev,
            "waterfall": scenario_waterfall(base_ev, layer_totals, model.other_claims),
        },
        {
            "label": "Healthy",
            "enterprise_value": high_ev,
            "waterfall": scenario_waterfall(high_ev, layer_totals, model.other_claims),
        },
    ]

    return {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "workbook": {
            "name": model.workbook_name,
            "base_ebitda": model.base_ebitda,
            "multiple_low": model.multiple_low,
            "multiple_mid": model.multiple_mid,
            "multiple_high": model.multiple_high,
            "listed_debt": model.listed_debt,
            "other_claims": model.other_claims,
            "total_debt_to_clear": model.total_debt_to_clear,
            "instruments": [asdict(instrument) for instrument in model.instruments],
        },
        "assumptions": asdict(assumptions),
        "summary": {
            "insolvency_probability": insolvency_probability,
            "positive_equity_probability": float(np.mean(equity_value > 0)),
            "expected_equity_value": float(np.mean(equity_value)),
            "equity_distribution": summarize_distribution(equity_value),
            "enterprise_value_distribution": summarize_distribution(enterprise_value),
            "ebitda_distribution": summarize_distribution(ebitda),
            "exit_multiple_distribution": summarize_distribution(exit_multiple),
        },
        "charts": {
            "enterprise_value_histogram": build_histogram(enterprise_value, bins=18),
            "equity_histogram": build_histogram(equity_value, bins=14),
        },
        "layer_stats": layer_stats,
        "instrument_stats": instrument_stats,
        "scenarios": scenarios,
        "sensitivity": build_sensitivity_matrix(model),
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def json_for_html(payload: Any) -> str:
    return json.dumps(payload, separators=(",", ":"), ensure_ascii=True).replace("</", "<\\/")


def render_report(report: dict[str, Any]) -> str:
    data_blob = json_for_html(report)
    title = escape(report["workbook"]["name"])
    generated_at = escape(report["generated_at"])
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>TSE Cap Table Simulation</title>
<style>
*,*::before,*::after{{box-sizing:border-box}}
:root{{
  --bg:#f6f4ef;
  --paper:#ffffff;
  --paper-strong:#fcfbf8;
  --border:#ddd6ca;
  --ink:#1e1b16;
  --muted:#746b5f;
  --accent:#20262d;
  --accent-soft:#e8ecef;
  --good:#246a4a;
  --warn:#a4631b;
  --bad:#8b3d35;
  --shadow:0 22px 60px rgba(28, 24, 18, 0.08);
}}
html{{
  background:
    radial-gradient(circle at top left, rgba(232,236,239,.9), transparent 28rem),
    linear-gradient(180deg, #fbfaf8 0%, #f3efe7 100%);
  min-height:100%;
}}
body{{
  margin:0;
  color:var(--ink);
  font-family:"Iowan Old Style","Palatino Linotype","Book Antiqua",Georgia,serif;
  min-height:100vh;
}}
.shell{{max-width:1320px;margin:0 auto;padding:32px 22px 56px}}
.hero{{
  background:linear-gradient(180deg, rgba(255,255,255,.95), rgba(251,249,245,.98));
  border:1px solid rgba(221,214,202,.9);
  border-radius:28px;
  box-shadow:var(--shadow);
  padding:32px;
  position:relative;
  overflow:hidden;
}}
.hero::after{{
  content:"";
  position:absolute;
  inset:auto -80px -90px auto;
  width:280px;
  height:280px;
  border-radius:50%;
  background:radial-gradient(circle, rgba(232,236,239,.95), rgba(232,236,239,0));
}}
.eyebrow,.label,.small,.table th,.table td,.axis-label,.cell-note{{
  font-family:"SFMono-Regular","IBM Plex Mono","Menlo",monospace;
}}
.eyebrow{{font-size:11px;letter-spacing:.16em;text-transform:uppercase;color:var(--muted);margin-bottom:16px}}
.hero h1{{font-size:clamp(2.3rem, 5vw, 4.6rem);line-height:.95;letter-spacing:-.05em;margin:0;max-width:9ch}}
.hero p{{max-width:48rem;font-size:1.02rem;line-height:1.7;color:#453d33;margin:18px 0 0}}
.hero-grid{{display:grid;grid-template-columns:1.6fr .9fr;gap:22px;align-items:end}}
.meta-card{{background:rgba(250,248,244,.92);border:1px solid var(--border);border-radius:22px;padding:18px 18px 16px;position:relative;z-index:1}}
.meta-line{{display:flex;justify-content:space-between;gap:14px;padding:10px 0;border-bottom:1px solid rgba(221,214,202,.8)}}
.meta-line:last-child{{border-bottom:none;padding-bottom:0}}
.meta-key{{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.12em}}
.meta-value{{font-size:13px;text-align:right}}
.grid{{display:grid;gap:18px;margin-top:18px}}
.metrics{{grid-template-columns:repeat(4,minmax(0,1fr))}}
.card{{
  background:rgba(255,255,255,.92);
  border:1px solid rgba(221,214,202,.92);
  border-radius:24px;
  padding:22px;
  box-shadow:0 10px 28px rgba(28,24,18,.04);
}}
.metric-value{{font-size:clamp(1.8rem,4vw,3rem);line-height:.95;letter-spacing:-.05em}}
.metric-value.good{{color:var(--good)}}
.metric-value.warn{{color:var(--warn)}}
.metric-value.bad{{color:var(--bad)}}
.label{{font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);margin-top:8px}}
.section{{margin-top:18px}}
.split{{display:grid;grid-template-columns:1.08fr .92fr;gap:18px}}
.histogram{{display:flex;align-items:flex-end;gap:6px;height:280px;margin-top:18px;padding-top:10px}}
.bar{{flex:1;border-radius:14px 14px 4px 4px;background:linear-gradient(180deg, #20262d, #cfd7dc);position:relative;min-width:12px}}
.bar::after{{content:attr(data-tip);position:absolute;left:50%;bottom:calc(100% + 10px);transform:translateX(-50%);background:#1f252c;color:#fff;padding:7px 8px;border-radius:8px;font-size:10px;line-height:1.4;opacity:0;pointer-events:none;transition:opacity .15s}}
.bar:hover::after{{opacity:1}}
.axis{{display:flex;justify-content:space-between;margin-top:12px;color:var(--muted);font-size:11px}}
.recovery-list{{display:grid;gap:12px;margin-top:16px}}
.recovery-row{{display:grid;grid-template-columns:140px 1fr 110px;gap:16px;align-items:center;padding:14px 16px;border:1px solid rgba(221,214,202,.9);border-radius:18px;background:linear-gradient(180deg,#fff,#fcfaf6)}}
.recovery-name{{font-size:1rem}}
.track{{height:14px;background:#efe9de;border-radius:999px;overflow:hidden;position:relative}}
.fill{{height:100%;border-radius:999px;background:linear-gradient(90deg,#20262d,#5d7688)}}
.recovery-meta{{text-align:right}}
.kpi-inline{{display:flex;gap:20px;flex-wrap:wrap;margin-top:14px;color:#4f483f}}
.kpi-inline span{{display:block}}
.kpi-inline small{{font-family:"SFMono-Regular","IBM Plex Mono","Menlo",monospace;color:var(--muted);font-size:10px;text-transform:uppercase;letter-spacing:.1em}}
.scenario-grid{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:16px}}
.scenario-card{{padding:18px;border-radius:22px;border:1px solid rgba(221,214,202,.9);background:linear-gradient(180deg,#ffffff,#fbfaf7)}}
.scenario-card h3{{margin:0 0 8px;font-size:1.35rem;letter-spacing:-.03em}}
.stack{{display:grid;gap:10px;margin-top:14px}}
.stack-row{{display:grid;grid-template-columns:1fr auto;gap:14px;padding-top:10px;border-top:1px solid rgba(221,214,202,.85)}}
.stack-row:first-child{{border-top:none;padding-top:0}}
.table-wrap{{overflow:auto;margin-top:14px}}
.table{{width:100%;border-collapse:collapse}}
.table th{{font-size:10px;text-transform:uppercase;letter-spacing:.14em;color:var(--muted);text-align:left;padding:0 0 10px}}
.table td{{padding:14px 0;border-top:1px solid rgba(221,214,202,.85);font-size:13px;vertical-align:top}}
.table td strong{{display:block;font-size:14px;margin-bottom:3px}}
.heatmap{{display:grid;grid-template-columns:120px repeat(5,1fr);gap:8px;margin-top:18px;align-items:stretch}}
.heatmap-header{{font-size:11px;color:var(--muted);padding:10px 8px}}
.heatmap-label{{display:flex;align-items:center;font-size:11px;color:var(--muted);padding-left:4px}}
.cell{{padding:12px;border-radius:16px;border:1px solid rgba(221,214,202,.85);min-height:96px;display:flex;flex-direction:column;justify-content:space-between}}
.cell-note{{font-size:10px;line-height:1.5}}
.footer{{margin-top:18px;color:var(--muted);font-size:12px;text-align:right}}
@media (max-width: 1024px){{
  .hero-grid,.split,.scenario-grid,.metrics{{grid-template-columns:1fr 1fr}}
}}
@media (max-width: 760px){{
  .shell{{padding:18px 14px 40px}}
  .hero,.card{{padding:20px}}
  .hero-grid,.split,.scenario-grid,.metrics{{grid-template-columns:1fr}}
  .recovery-row{{grid-template-columns:1fr;gap:10px}}
  .recovery-meta{{text-align:left}}
  .heatmap{{grid-template-columns:1fr}}
  .heatmap-header{{display:none}}
  .heatmap-label{{padding-left:0}}
}}
</style>
</head>
<body>
<div class="shell">
  <section class="hero">
    <div class="hero-grid">
      <div>
        <div class="eyebrow">TSE cap table recovery engine</div>
        <h1>Monte Carlo exit waterfall.</h1>
        <p>
          A standalone view over <strong>{title}</strong> that simulates exit enterprise value, applies the capital stack in priority order,
          and surfaces insolvency odds, debt coverage, recovery ranges, and residual equity.
        </p>
      </div>
      <div class="meta-card">
        <div class="meta-line"><div class="meta-key">Generated</div><div class="meta-value">{generated_at}</div></div>
        <div class="meta-line"><div class="meta-key">Runs</div><div class="meta-value" id="runs"></div></div>
        <div class="meta-line"><div class="meta-key">EBITDA Base</div><div class="meta-value" id="ebitda-base"></div></div>
        <div class="meta-line"><div class="meta-key">Multiple Range</div><div class="meta-value" id="multiple-range"></div></div>
        <div class="meta-line"><div class="meta-key">Debt To Clear</div><div class="meta-value" id="debt-total"></div></div>
      </div>
    </div>
  </section>

  <section class="grid metrics">
    <article class="card">
      <div class="metric-value bad" id="insolvency-prob"></div>
      <div class="label">Probability EV fails to clear debt</div>
    </article>
    <article class="card">
      <div class="metric-value" id="median-ev"></div>
      <div class="label">Median enterprise value</div>
    </article>
    <article class="card">
      <div class="metric-value good" id="equity-prob"></div>
      <div class="label">Probability equity has value</div>
    </article>
    <article class="card">
      <div class="metric-value warn" id="expected-equity"></div>
      <div class="label">Expected residual equity</div>
    </article>
  </section>

  <section class="grid split section">
    <article class="card">
      <div class="eyebrow">Enterprise value distribution</div>
      <div class="metric-value" id="mean-ev"></div>
      <div class="label">Mean simulated EV</div>
      <div class="kpi-inline">
        <span><strong id="ev-p10"></strong><br><small>P10</small></span>
        <span><strong id="ev-p50"></strong><br><small>P50</small></span>
        <span><strong id="ev-p90"></strong><br><small>P90</small></span>
      </div>
      <div class="histogram" id="ev-histogram"></div>
      <div class="axis">
        <div class="axis-label" id="ev-axis-start"></div>
        <div class="axis-label" id="ev-axis-end"></div>
      </div>
    </article>
    <article class="card">
      <div class="eyebrow">Layer recovery outlook</div>
      <div class="recovery-list" id="layer-recoveries"></div>
    </article>
  </section>

  <section class="grid split section">
    <article class="card">
      <div class="eyebrow">Deterministic reference cases</div>
      <div class="scenario-grid" id="scenario-grid"></div>
    </article>
    <article class="card">
      <div class="eyebrow">Sensitivity map</div>
      <div class="label">Each cell shows EV, 2L recovery, and residual equity.</div>
      <div class="heatmap" id="heatmap"></div>
    </article>
  </section>

  <section class="card section">
    <div class="eyebrow">Instrument detail</div>
    <div class="table-wrap">
      <table class="table">
        <thead>
          <tr>
            <th>Instrument</th>
            <th>Layer</th>
            <th>Principal</th>
            <th>Maturity</th>
            <th>Interest</th>
            <th>Expected Recovery</th>
            <th>Money-Good</th>
          </tr>
        </thead>
        <tbody id="instrument-table"></tbody>
      </table>
    </div>
  </section>

  <div class="footer">Standalone HTML dashboard generated by <code>simulate_tse_cap_table.py</code>.</div>
</div>

<script id="report-data" type="application/json">{data_blob}</script>
<script>
const report = JSON.parse(document.getElementById('report-data').textContent);

const money = value => `$${{Number(value).toLocaleString(undefined, {{minimumFractionDigits: 1, maximumFractionDigits: 1}})}}M`;
const pct = value => `${{(Number(value) * 100).toFixed(1)}}%`;
const clamp = (value, min, max) => Math.min(max, Math.max(min, value));

document.getElementById('runs').textContent = report.assumptions.runs.toLocaleString();
document.getElementById('ebitda-base').textContent = money(report.workbook.base_ebitda);
document.getElementById('multiple-range').textContent = `${{report.workbook.multiple_low.toFixed(1)}}x to ${{report.workbook.multiple_high.toFixed(1)}}x`;
document.getElementById('debt-total').textContent = money(report.workbook.total_debt_to_clear);

document.getElementById('insolvency-prob').textContent = pct(report.summary.insolvency_probability);
document.getElementById('median-ev').textContent = money(report.summary.enterprise_value_distribution.p50);
document.getElementById('equity-prob').textContent = pct(report.summary.positive_equity_probability);
document.getElementById('expected-equity').textContent = money(report.summary.expected_equity_value);

document.getElementById('mean-ev').textContent = money(report.summary.enterprise_value_distribution.mean);
document.getElementById('ev-p10').textContent = money(report.summary.enterprise_value_distribution.p10);
document.getElementById('ev-p50').textContent = money(report.summary.enterprise_value_distribution.p50);
document.getElementById('ev-p90').textContent = money(report.summary.enterprise_value_distribution.p90);

const evHistogram = report.charts.enterprise_value_histogram;
const maxCount = Math.max(...evHistogram.map(bin => bin.count), 1);
const histEl = document.getElementById('ev-histogram');
evHistogram.forEach(bin => {{
  const bar = document.createElement('div');
  bar.className = 'bar';
  bar.style.height = `${{Math.max(8, (bin.count / maxCount) * 100)}}%`;
  bar.dataset.tip = `${{money(bin.start)}} to ${{money(bin.end)}} · ${{bin.count.toLocaleString()}} runs`;
  histEl.appendChild(bar);
}});
document.getElementById('ev-axis-start').textContent = money(evHistogram[0].start);
document.getElementById('ev-axis-end').textContent = money(evHistogram[evHistogram.length - 1].end);

const layerRecoveries = document.getElementById('layer-recoveries');
report.layer_stats.forEach(layer => {{
  const row = document.createElement('div');
  row.className = 'recovery-row';
  row.innerHTML = `
    <div>
      <div class="label">${{layer.layer}}</div>
      <div class="recovery-name">${{money(layer.current_claim)}}</div>
    </div>
    <div>
      <div class="track"><div class="fill" style="width:${{clamp(layer.expected_recovery_pct * 100, 0, 100)}}%"></div></div>
      <div class="kpi-inline">
        <span><strong>${{pct(layer.recovery_distribution.p10)}}</strong><br><small>P10</small></span>
        <span><strong>${{pct(layer.recovery_distribution.p50)}}</strong><br><small>P50</small></span>
        <span><strong>${{pct(layer.recovery_distribution.p90)}}</strong><br><small>P90</small></span>
      </div>
    </div>
    <div class="recovery-meta">
      <div><strong>${{pct(layer.expected_recovery_pct)}}</strong></div>
      <div class="label">Expected recovery</div>
      <div style="margin-top:8px"><strong>${{pct(layer.money_good_probability)}}</strong></div>
      <div class="label">Money-good</div>
    </div>
  `;
  layerRecoveries.appendChild(row);
}});

const scenarioGrid = document.getElementById('scenario-grid');
report.scenarios.forEach(scenario => {{
  const card = document.createElement('div');
  card.className = 'scenario-card';
  const rows = scenario.waterfall.map(item => `
    <div class="stack-row">
      <div>${{item.layer}}</div>
      <div>${{item.layer === 'Equity residual' ? money(item.recovered) : `${{money(item.recovered)}} · ${{pct(item.recovery_pct)}}`}}</div>
    </div>
  `).join('');
  card.innerHTML = `
    <h3>${{scenario.label}}</h3>
    <div class="label">${{money(scenario.enterprise_value)}} enterprise value</div>
    <div class="stack">${{rows}}</div>
  `;
  scenarioGrid.appendChild(card);
}});

const heatmap = document.getElementById('heatmap');
heatmap.innerHTML = '<div></div>' + report.sensitivity.multiple_values.map(value => `<div class="heatmap-header">${{value.toFixed(2)}}x</div>`).join('');
const maxEquity = report.sensitivity.max_equity || 1;
report.sensitivity.cells.forEach((row, rowIndex) => {{
  const label = document.createElement('div');
  label.className = 'heatmap-label';
  label.textContent = money(report.sensitivity.ebitda_values[rowIndex]) + ' EBITDA';
  heatmap.appendChild(label);
  row.forEach(cell => {{
    const tile = document.createElement('div');
    tile.className = 'cell';
    const intensity = clamp(cell.equity_value / maxEquity, 0, 1);
    const ink = intensity > 0.56 ? '#ffffff' : '#1f1c17';
    tile.style.background = `linear-gradient(180deg, rgba(32,38,45,${{0.08 + intensity * 0.86}}), rgba(232,236,239,0.88))`;
    tile.style.color = ink;
    tile.innerHTML = `
      <div><strong>${{money(cell.enterprise_value)}}</strong></div>
      <div class="cell-note" style="color:${{ink === '#ffffff' ? 'rgba(255,255,255,0.84)' : '#5c5348'}}">
        2L recovery: ${{pct(cell.second_lien_recovery)}}<br>
        Equity: ${{money(cell.equity_value)}}
      </div>
    `;
    heatmap.appendChild(tile);
  }});
}});

const instrumentTable = document.getElementById('instrument-table');
report.instrument_stats.forEach(item => {{
  const row = document.createElement('tr');
  row.innerHTML = `
    <td><strong>${{item.instrument}}</strong></td>
    <td>${{item.layer}}</td>
    <td>${{money(item.principal)}}</td>
    <td>${{item.maturity}}</td>
    <td>${{item.interest_type}}</td>
    <td>${{money(item.expected_recovered)}} · ${{pct(item.expected_recovery_pct)}}</td>
    <td>${{pct(item.money_good_probability)}}</td>
  `;
  instrumentTable.appendChild(row);
}});
</script>
</body>
</html>
"""


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.input)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    model = load_cap_table(workbook_path)
    assumptions = SimulationAssumptions(
        runs=args.runs,
        seed=args.seed,
        ebitda_cv=args.ebitda_cv,
        debt_volatility=args.debt_volatility,
    )
    report = simulate(model, assumptions)

    json_path = output_dir / "tse_cap_table_simulation.json"
    csv_layers_path = output_dir / "tse_cap_table_layer_stats.csv"
    csv_instruments_path = output_dir / "tse_cap_table_instrument_stats.csv"
    csv_scenarios_path = output_dir / "tse_cap_table_scenarios.csv"
    html_path = output_dir / "tse_cap_table_report.html"

    json_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    write_csv(
        csv_layers_path,
        [
            {
                "layer": row["layer"],
                "current_claim_m": f"{row['current_claim']:.4f}",
                "expected_claim_at_exit_m": f"{row['expected_claim_at_exit']:.4f}",
                "expected_recovered_m": f"{row['expected_recovered']:.4f}",
                "expected_recovery_pct": f"{row['expected_recovery_pct']:.6f}",
                "money_good_probability": f"{row['money_good_probability']:.6f}",
                "p10_recovery_pct": f"{row['recovery_distribution']['p10']:.6f}",
                "p50_recovery_pct": f"{row['recovery_distribution']['p50']:.6f}",
                "p90_recovery_pct": f"{row['recovery_distribution']['p90']:.6f}",
            }
            for row in report["layer_stats"]
        ],
        [
            "layer",
            "current_claim_m",
            "expected_claim_at_exit_m",
            "expected_recovered_m",
            "expected_recovery_pct",
            "money_good_probability",
            "p10_recovery_pct",
            "p50_recovery_pct",
            "p90_recovery_pct",
        ],
    )
    write_csv(
        csv_instruments_path,
        [
            {
                "layer": row["layer"],
                "instrument": row["instrument"],
                "principal_m": f"{row['principal']:.4f}",
                "maturity": row["maturity"],
                "interest_type": row["interest_type"],
                "expected_claim_at_exit_m": f"{row['expected_claim_at_exit']:.4f}",
                "expected_recovered_m": f"{row['expected_recovered']:.4f}",
                "expected_recovery_pct": f"{row['expected_recovery_pct']:.6f}",
                "money_good_probability": f"{row['money_good_probability']:.6f}",
            }
            for row in report["instrument_stats"]
        ],
        [
            "layer",
            "instrument",
            "principal_m",
            "maturity",
            "interest_type",
            "expected_claim_at_exit_m",
            "expected_recovered_m",
            "expected_recovery_pct",
            "money_good_probability",
        ],
    )
    write_csv(
        csv_scenarios_path,
        [
            {
                "scenario": scenario["label"],
                "enterprise_value_m": f"{scenario['enterprise_value']:.4f}",
                "layer": row["layer"],
                "claim_m": f"{row['claim']:.4f}",
                "recovered_m": f"{row['recovered']:.4f}",
                "recovery_pct": f"{row['recovery_pct']:.6f}",
            }
            for scenario in report["scenarios"]
            for row in scenario["waterfall"]
        ],
        [
            "scenario",
            "enterprise_value_m",
            "layer",
            "claim_m",
            "recovered_m",
            "recovery_pct",
        ],
    )
    html_path.write_text(render_report(report), encoding="utf-8")

    print(f"Wrote {json_path}")
    print(f"Wrote {csv_layers_path}")
    print(f"Wrote {csv_instruments_path}")
    print(f"Wrote {csv_scenarios_path}")
    print(f"Wrote {html_path}")


if __name__ == "__main__":
    main()
